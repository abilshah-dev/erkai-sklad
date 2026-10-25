"""Загрузка номенклатуры, цен и остатков из Excel-выгрузки 1С.

Как выгрузить из 1С: отчёт «Остатки товаров» (или прайс-лист) → Сохранить как Excel (.xlsx).
Нужны колонки (порядок не важен, ищем по названию заголовка):
  Код | Наименование | Группа | Ед. изм. | Цена | Остаток
Колонки «Группа» и «Остаток» необязательны.

Запуск на сервере (из папки проекта, с активным venv):
  python -m tools.import_1c путь/к/файлу.xlsx
  python -m tools.import_1c файл.xlsx --no-stock     # обновить только товары и цены
"""
from __future__ import annotations

import sys

from openpyxl import load_workbook

from app import db

HEADERS = {
    "code": ("код", "артикул"),
    "name": ("наименование", "номенклатура", "товар"),
    "group": ("группа", "категория", "родитель"),
    "unit": ("ед", "единица"),
    "price": ("цена",),
    "stock": ("остаток", "количество", "кол-во"),
}


def find_columns(row) -> dict:
    cols = {}
    for idx, cell in enumerate(row):
        title = str(cell or "").strip().lower()
        for key, variants in HEADERS.items():
            if key not in cols and any(title.startswith(v) for v in variants):
                cols[key] = idx
    return cols


def read_rows(path: str, with_stock: bool) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    # Ищем строку заголовков в первых 15 строках (в отчётах 1С сверху бывает шапка)
    for h, row in enumerate(rows[:15]):
        cols = find_columns(row)
        if "code" in cols and "name" in cols:
            break
    else:
        sys.exit("Не нашёл заголовки «Код» и «Наименование» в первых 15 строках.")

    out = []
    for row in rows[h + 1:]:
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
        code, name = get("code"), get("name")
        if not code or not name:
            continue  # пустые и итоговые строки пропускаем
        item = {
            "code": str(code).strip(),
            "name": str(name).strip(),
            "group": str(get("group") or "").strip(),
            "unit": str(get("unit") or "шт").strip(),
            "price": float(get("price") or 0),
        }
        if with_stock and "stock" in cols:
            item["stock"] = max(0, round(float(get("stock") or 0)))
        out.append(item)
    return out


async def main(path: str, with_stock: bool):
    rows = read_rows(path, with_stock)
    print(f"Прочитано строк: {len(rows)}")
    for i in range(0, len(rows), 200):  # отправляем пачками по 200
        res = await db.rpc("import_catalog", {"p_rows": rows[i:i + 200]})
        print(f"Загружено: {i + res['rows']}")
    print("Готово.")


if __name__ == "__main__":
    import asyncio

    if len(sys.argv) < 2:
        sys.exit(__doc__)
    asyncio.run(main(sys.argv[1], "--no-stock" not in sys.argv))
